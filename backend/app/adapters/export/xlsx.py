"""Exporting an analysis to Excel (spec §34 criteria 11 and 12).

The export is the deliverable a preparer actually hands to a reviewer, so two
things matter more than layout:

**The figures must be the figures.** openpyxl converts ``Decimal`` to a binary
float on write, which silently loses precision beyond about sixteen significant
digits. Rather than accept that, or give up numeric cells and write strings —
which would break every sum a reviewer tries — each monetary value is checked to
round-trip exactly, and an export that would lose a digit fails instead of
shipping a wrong number. On realistic statements the check never fires.

**A file that did not reconcile must never be mistakable for one that did.**
Spec §19 permits taking unreconciled output away to investigate it, but every
sheet is then watermarked, and the reconciliation sheet states exactly which
check failed and by how much.
"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.core.product import (
    DISCLAIMER_EN,
    DISCLAIMER_KO,
    LIMITATIONS_EN,
    UNRECONCILED_WATERMARK,
)
from app.domain.classification import ClassificationDecision
from app.domain.enums import DecompositionStatus
from app.domain.extraction import ExtractedStatement
from app.domain.impact import ImpactAnalysis, StepKind, Unit
from app.domain.statement import ClassifiedLine, Ifrs18Statement
from app.domain.validation import ValidationReport

AMOUNT_FORMAT = "#,##0"
RATIO_FORMAT = "#,##0.0000"

_HEADER_FILL = PatternFill("solid", fgColor="E8E8E4")
_WATERMARK_FILL = PatternFill("solid", fgColor="F6D5D0")
_HEADER_FONT = Font(bold=True)
_TITLE_FONT = Font(bold=True, size=13)
_WATERMARK_FONT = Font(bold=True, color="9C3328")


class ExportPrecisionError(ValueError):
    """A figure cannot be written to Excel without losing precision."""


@dataclass(frozen=True, slots=True)
class ExportPackage:
    """Everything an export needs, gathered once."""

    source: ExtractedStatement
    classified: tuple[ClassifiedLine, ...]
    reconstructed: Ifrs18Statement
    validation: ValidationReport
    impact: ImpactAnalysis
    project_name: str = "IFRS 18 analysis"
    rule_set_version: str | None = None
    catalog_version: str | None = None

    @property
    def reconciled(self) -> bool:
        return self.validation.passed


def excel_number(value: Decimal) -> float:
    """Convert a figure for Excel, refusing to lose a digit.

    ``repr`` of a double is the shortest decimal that round-trips it, so
    comparing against that detects exactly the cases where the conversion is
    lossy — and only those.
    """
    try:
        as_float = float(value)
        if Decimal(repr(as_float)) == value:
            return as_float
    except (OverflowError, ValueError) as exc:  # pragma: no cover - extreme input
        raise ExportPrecisionError(f"{value} cannot be represented in Excel") from exc

    raise ExportPrecisionError(
        f"{value} has more precision than a spreadsheet cell can hold, so "
        "exporting it would silently change the figure. Reduce the "
        "presentation scale, or export the audit trail instead."
    )


# ---------------------------------------------------------------------------
# Sheet helpers
# ---------------------------------------------------------------------------


def _write_row(
    sheet: Worksheet,
    row: int,
    values: list[Any],
    *,
    bold: bool = False,
    number_format: str = AMOUNT_FORMAT,
) -> int:
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=column)
        if isinstance(value, Decimal):
            cell.value = excel_number(value)
            cell.number_format = number_format
        else:
            cell.value = value
        if bold:
            cell.font = _HEADER_FONT
    return row + 1


def _write_header(sheet: Worksheet, row: int, headers: list[str]) -> int:
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=column, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    # Keep the header visible; a financial sheet is scrolled, not read once.
    sheet.freeze_panes = f"A{row + 1}"
    return row + 1


def _start_sheet(workbook: Workbook, title: str, package: ExportPackage) -> tuple[Worksheet, int]:
    """Create a sheet, watermarking it when the analysis did not reconcile."""
    sheet = workbook.create_sheet(title)
    row = 1
    if not package.reconciled:
        cell = sheet.cell(row=row, column=1, value=UNRECONCILED_WATERMARK)
        cell.font = _WATERMARK_FONT
        cell.fill = _WATERMARK_FILL
        row += 2
    return sheet, row


def _autosize(sheet: Worksheet, widths: dict[int, int]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width


# ---------------------------------------------------------------------------
# Sheets
# ---------------------------------------------------------------------------


def _summary_sheet(workbook: Workbook, package: ExportPackage) -> None:
    sheet, row = _start_sheet(workbook, "Summary", package)

    sheet.cell(row=row, column=1, value=package.project_name).font = _TITLE_FONT
    row += 2

    row = _write_row(
        sheet,
        row,
        [
            "Reconciliation",
            "PASSED" if package.reconciled else "FAILED",
        ],
        bold=True,
    )
    row = _write_row(sheet, row, ["Currency", package.reconstructed.currency])
    row = _write_row(sheet, row, ["Presentation scale (power of ten)", package.reconstructed.scale])
    if package.rule_set_version:
        row = _write_row(sheet, row, ["IFRS 18 rule set", package.rule_set_version])
    if package.catalog_version:
        row = _write_row(sheet, row, ["Account catalog", package.catalog_version])
    row += 1

    row = _write_header(
        sheet, row, ["Measure", "Before", "After", "Change", "Change %", "Change (bp)"]
    )
    for kpi in package.impact.kpis:
        number_format = RATIO_FORMAT if kpi.unit is Unit.PERCENT else AMOUNT_FORMAT
        row = _write_row(
            sheet,
            row,
            [
                kpi.key,
                kpi.before if kpi.before is not None else "not applicable",
                kpi.after,
                kpi.change if kpi.change is not None else "—",
                kpi.change_pct if kpi.change_pct is not None else "—",
                kpi.change_bps if kpi.change_bps is not None else "—",
            ],
            number_format=number_format,
        )

    row += 1
    sheet.cell(row=row, column=1, value="A measure shown as 'not applicable' did not")
    row += 1
    sheet.cell(
        row=row,
        column=1,
        value="exist under the entity's previous presentation; IFRS 18 introduced it.",
    )
    row += 2

    row = _write_row(sheet, row, ["Disclaimer"], bold=True)
    row = _write_row(sheet, row, [DISCLAIMER_EN])
    row = _write_row(sheet, row, [DISCLAIMER_KO])
    row += 1
    row = _write_row(sheet, row, ["Scope limitations"], bold=True)
    for limitation in LIMITATIONS_EN:
        row = _write_row(sheet, row, [f"— {limitation}"])

    _autosize(sheet, {1: 46, 2: 18, 3: 18, 4: 16, 5: 14, 6: 14})


def _statement_sheet(workbook: Workbook, package: ExportPackage) -> None:
    sheet, row = _start_sheet(workbook, "IFRS 18 Statement", package)
    row = _write_header(sheet, row, ["Category", "Account", "Amount"])

    subtotals = {s.key: s for s in package.reconstructed.subtotals}
    for section in package.reconstructed.sections:
        if section.is_empty:
            continue
        row = _write_row(sheet, row, [section.label_en], bold=True)
        for item in section.lines:
            row = _write_row(sheet, row, ["", item.line.raw_label, item.amount])
        row = _write_row(sheet, row, ["", f"{section.label_en} total", section.total], bold=True)
        row += 1

    row = _write_row(sheet, row, ["Subtotals"], bold=True)
    for subtotal in subtotals.values():
        label = subtotal.label_en
        if not subtotal.presented:
            # Never simply omitted: a reader must see that the standard forbids
            # it, not be left to wonder why it is missing.
            label += "  [NOT PRESENTED]"
        row = _write_row(sheet, row, ["", label, subtotal.amount], bold=True)
        if subtotal.suppressed_reason:
            row = _write_row(sheet, row, ["", subtotal.suppressed_reason])

    _autosize(sheet, {1: 26, 2: 46, 3: 20})


def _impact_sheet(workbook: Workbook, package: ExportPackage) -> None:
    sheet, row = _start_sheet(workbook, "Impact", package)

    if not package.impact.comparable:
        row = _write_row(
            sheet,
            row,
            ["The source printed no operating subtotal, so there is nothing to compare against."],
            bold=True,
        )
        _autosize(sheet, {1: 90})
        return

    row = _write_row(sheet, row, ["Operating profit bridge"], bold=True)
    row = _write_header(sheet, row, ["Step", "Description", "Amount", "IFRS 18 category"])
    for step in package.impact.waterfall:
        row = _write_row(
            sheet,
            row,
            [
                step.kind.value,
                step.label_en,
                step.value,
                step.to_category.value if step.to_category else "",
            ],
            bold=step.kind in (StepKind.START, StepKind.END),
        )

    row += 2
    row = _write_row(sheet, row, ["Reclassifications, largest first"], bold=True)
    row = _write_header(
        sheet,
        row,
        [
            "Account",
            "Amount",
            "Previously",
            "IFRS 18 category",
            "Effect on operating profit",
            "Rule",
        ],
    )
    for item in package.impact.top_reclassifications:
        row = _write_row(
            sheet,
            row,
            [
                item.label,
                item.amount,
                item.reported_placement.value,
                item.ifrs18_category.value,
                item.impact_on_operating_profit,
                item.rule_id or "",
            ],
        )

    _autosize(sheet, {1: 30, 2: 18, 3: 22, 4: 22, 5: 26, 6: 24})


def _inclusion(item: ClassifiedLine) -> str:
    """Why a row does or does not contribute to the totals.

    An aggregate caption appears alongside the components it was broken into,
    which is what an audit trail is for — but a reader who cannot tell which
    rows sum could double count them, so each row says so outright.
    """
    if item.line.is_subtotal:
        return "no — subtotal printed by the source"
    if item.line.decomposition_status is DecompositionStatus.DECOMPOSED:
        return "no — decomposed; its components carry the amount"
    return "yes"


def _accounts_sheet(workbook: Workbook, package: ExportPackage) -> None:
    """Every line, with the source cell it came from (spec §18)."""
    sheet, row = _start_sheet(workbook, "Accounts", package)
    row = _write_header(
        sheet,
        row,
        [
            "Account (as printed)",
            "Amount",
            "Included in totals",
            "Normalized account",
            "IFRS 18 category",
            "Method",
            "Source file",
            "Sheet",
            "Cell",
            "Notes",
        ],
    )

    for item in package.classified:
        locator = item.line.locator
        included = item.is_summable
        row = _write_row(
            sheet,
            row,
            [
                item.line.raw_label,
                item.amount,
                _inclusion(item),
                item.normalized_account_code or "",
                item.category.value if included else "",
                item.decision.method.value if included else "",
                locator.source_file,
                locator.sheet or "",
                locator.cell or "",
                ", ".join(item.line.note_references),
            ],
        )

    row += 1
    row = _write_row(
        sheet,
        row,
        ["Total of rows included above", _summable_total(package)],
        bold=True,
    )

    _autosize(
        sheet,
        {1: 32, 2: 18, 3: 42, 4: 30, 5: 22, 6: 20, 7: 22, 8: 18, 9: 10, 10: 20},
    )


def _summable_total(package: ExportPackage) -> Decimal:
    return sum(
        (item.amount for item in package.classified if item.is_summable),
        start=Decimal(0),
    )


def _audit_sheet(workbook: Workbook, package: ExportPackage) -> None:
    """Spec §34 criterion 12: every classification carries its basis."""
    sheet, row = _start_sheet(workbook, "Audit trail", package)
    row = _write_header(
        sheet,
        row,
        [
            "Account",
            "IFRS 18 category",
            "Method",
            "Rule",
            "Citation",
            "Confidence",
            "Band",
            "Needs review",
            "Reasoning",
        ],
    )

    for item in package.classified:
        decision: ClassificationDecision = item.decision
        citations = "; ".join(evidence.reference for evidence in decision.evidence)
        row = _write_row(
            sheet,
            row,
            [
                item.line.raw_label,
                item.category.value,
                decision.method.value,
                decision.rule_id or "",
                citations,
                decision.confidence,
                decision.confidence_band.value,
                "YES" if decision.requires_human_review else "no",
                decision.reasoning or "",
            ],
            number_format=RATIO_FORMAT,
        )

    _autosize(sheet, {1: 30, 2: 22, 3: 20, 4: 24, 5: 34, 6: 12, 7: 10, 8: 14, 9: 60})


def _reconciliation_sheet(workbook: Workbook, package: ExportPackage) -> None:
    sheet, row = _start_sheet(workbook, "Reconciliation", package)
    row = _write_header(
        sheet,
        row,
        ["Check", "Result", "Severity", "Expected", "Actual", "Difference", "Tolerance", "Detail"],
    )

    for finding in package.validation.findings:
        row = _write_row(
            sheet,
            row,
            [
                finding.check,
                "PASS" if finding.passed else "FAIL",
                finding.severity.value,
                finding.expected if finding.expected is not None else "",
                finding.actual if finding.actual is not None else "",
                finding.delta if finding.delta is not None else "",
                finding.tolerance,
                finding.detail,
            ],
            bold=not finding.passed,
        )

    _autosize(sheet, {1: 34, 2: 10, 3: 12, 4: 18, 5: 18, 6: 16, 7: 12, 8: 80})


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------


def build_workbook(package: ExportPackage) -> Workbook:
    workbook = Workbook()
    default = workbook.active
    if default is not None:
        workbook.remove(default)

    _summary_sheet(workbook, package)
    _statement_sheet(workbook, package)
    _impact_sheet(workbook, package)
    _accounts_sheet(workbook, package)
    _audit_sheet(workbook, package)
    _reconciliation_sheet(workbook, package)
    return workbook


def export_to_path(package: ExportPackage, path: Path | str) -> Path:
    path = Path(path)
    build_workbook(package).save(path)
    return path
