"""Synthetic Korean income statements for testing extraction.

Generated rather than committed as binary, so the fixture is reviewable as code
and its arithmetic is visible. **These are synthetic.** They imitate the shape
of a K-IFRS 손익계산서 but are not drawn from any real filing, so they validate
the parser, not the account dictionary or the rule set. Replacing them with a
real anonymised statement remains an open task (see `docs/05-mvp-scope.md` §6).

Three variants exercise the conventions that actually differ between filings:

``PARENTHESES``  expenses printed as ``(70,000)`` — the common export style
``TRIANGLE``     expenses printed as ``△70,000`` — a Korean convention
``UNSIGNED``     every figure printed positive, sign left implicit

The third is the hard one: read as printed it fails the statement's own
arithmetic, which is what ``infer_signs_from_subtotals`` must resolve.
"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from enum import StrEnum
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from app.domain.enums import SubtotalKind


class SignStyle(StrEnum):
    PARENTHESES = "PARENTHESES"
    TRIANGLE = "TRIANGLE"
    UNSIGNED = "UNSIGNED"


@dataclass(frozen=True)
class Row:
    label: str
    #: Signed profit-or-loss effect — the truth the parser must recover.
    amount: Decimal
    depth: int = 0
    is_subtotal: bool = False
    subtotal_kind: SubtotalKind | None = None
    note: str | None = None


#: Figures in 백만원. The running totals are internally consistent:
#:   매출총이익  = 1,000,000 - 700,000            =  300,000
#:   영업이익    =   300,000 - 180,000            =  120,000
#:   법인세차감전순이익 = 120,000 + 12,000 + 8,000 + 5,000 - 14,000 - 9,000 = 122,000
#:   당기순이익  =   122,000 -  26,840            =   95,160
STATEMENT_ROWS: tuple[Row, ...] = (
    Row("매출액", Decimal("1000000"), note="주석 21"),
    Row("매출원가", Decimal("-700000"), note="주석 22"),
    Row("매출총이익", Decimal("300000"), is_subtotal=True, subtotal_kind=SubtotalKind.GROSS_PROFIT),
    Row("판매비와관리비", Decimal("-180000"), note="주석 23"),
    Row(
        "영업이익",
        Decimal("120000"),
        is_subtotal=True,
        subtotal_kind=SubtotalKind.REPORTED_OPERATING_PROFIT,
    ),
    Row("기타수익", Decimal("12000"), depth=1, note="주석 24"),
    Row("금융수익", Decimal("8000"), depth=1, note="주석 25"),
    Row("지분법이익", Decimal("5000"), depth=1, note="주석 13"),
    Row("금융비용", Decimal("-14000"), depth=1, note="주석 25"),
    Row("기타비용", Decimal("-9000"), depth=1, note="주석 24"),
    Row(
        "법인세차감전순이익",
        Decimal("122000"),
        is_subtotal=True,
        subtotal_kind=SubtotalKind.PROFIT_BEFORE_TAX,
    ),
    Row("법인세비용", Decimal("-26840"), note="주석 26"),
    Row(
        "당기순이익",
        Decimal("95160"),
        is_subtotal=True,
        subtotal_kind=SubtotalKind.PROFIT_FOR_THE_PERIOD,
    ),
)


def _render(amount: Decimal, style: SignStyle) -> str | Decimal:
    if amount >= 0:
        return amount
    magnitude = abs(amount)
    if style is SignStyle.PARENTHESES:
        return f"({magnitude:,})"
    if style is SignStyle.TRIANGLE:
        return f"△{magnitude:,}"
    return magnitude  # UNSIGNED: the sign is simply not printed


def build_workbook(
    path: Path,
    *,
    style: SignStyle = SignStyle.PARENTHESES,
    sheet_name: str = "손익계산서",
    include_comparative: bool = True,
) -> Path:
    """Write a synthetic statement and return its path."""
    workbook = Workbook()
    # A leading sheet of unrelated content, so the reader must actually locate
    # the statement rather than assume it is first.
    cover = workbook.active
    assert cover is not None
    cover.title = "표지"
    cover["A1"] = "연결재무제표"
    cover["A3"] = "제 55 기"

    sheet = workbook.create_sheet(sheet_name)
    sheet["A1"] = "연결 포괄손익계산서"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = "제55기 2025.01.01 부터 2025.12.31 까지"
    sheet["A3"] = "(단위: 백만원)"

    header_row = 5
    sheet.cell(row=header_row, column=1, value="과목")
    sheet.cell(row=header_row, column=2, value="주석")
    sheet.cell(row=header_row, column=3, value="제55기")
    if include_comparative:
        sheet.cell(row=header_row, column=4, value="제54기")
    for column in range(1, 5 if include_comparative else 4):
        sheet.cell(row=header_row, column=column).font = Font(bold=True)

    for offset, row in enumerate(STATEMENT_ROWS):
        excel_row = header_row + 1 + offset
        label_cell = sheet.cell(row=excel_row, column=1, value=row.label)
        label_cell.alignment = Alignment(indent=row.depth)
        if row.is_subtotal:
            label_cell.font = Font(bold=True)
        if row.note:
            sheet.cell(row=excel_row, column=2, value=row.note)
        sheet.cell(row=excel_row, column=3, value=_render(row.amount, style))
        if include_comparative:
            # Prior period: same shape, scaled down, so the reader must not
            # silently mix columns.
            prior = (row.amount * Decimal("0.9")).quantize(Decimal("1"))
            sheet.cell(row=excel_row, column=4, value=_render(prior, style))

    sheet.column_dimensions["A"].width = 28
    workbook.save(path)
    return path


def expected_amounts() -> dict[str, Decimal]:
    return {row.label: row.amount for row in STATEMENT_ROWS}
