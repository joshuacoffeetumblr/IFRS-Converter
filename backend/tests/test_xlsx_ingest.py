"""Reading a workbook end to end, against the synthetic Korean fixtures."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.adapters.ingest.xlsx import (
    ReadOptions,
    StatementNotFoundError,
    classify_subtotal,
    detect_scale,
    read_income_statement,
)
from app.domain.enums import SignNormalization, SubtotalKind
from app.domain.extraction import infer_signs_from_subtotals, reconcile_extraction
from tests.fixtures.korean_income_statement import (
    STATEMENT_ROWS,
    SignStyle,
    build_workbook,
    expected_amounts,
)


@pytest.fixture
def workbook_path(tmp_path: Path) -> Path:
    return build_workbook(tmp_path / "fs.xlsx", style=SignStyle.PARENTHESES)


# ---------------------------------------------------------------------------
# Caption classification
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("label", "kind"),
    [
        ("매출총이익", SubtotalKind.GROSS_PROFIT),
        ("영업이익", SubtotalKind.REPORTED_OPERATING_PROFIT),
        ("영업손실", SubtotalKind.REPORTED_OPERATING_PROFIT),
        ("법인세차감전순이익", SubtotalKind.PROFIT_BEFORE_TAX),
        ("법인세비용차감전순이익", SubtotalKind.PROFIT_BEFORE_TAX),
        ("당기순이익", SubtotalKind.PROFIT_FOR_THE_PERIOD),
        ("당 기 순 이 익", SubtotalKind.PROFIT_FOR_THE_PERIOD),
    ],
)
def test_subtotal_captions_are_recognised(label: str, kind: SubtotalKind) -> None:
    assert classify_subtotal(label) is kind


@pytest.mark.parametrize("label", ["매출액", "이자수익", "판매비와관리비", "지분법이익", ""])
def test_detail_captions_are_not_subtotals(label: str) -> None:
    assert classify_subtotal(label) is None


def test_profit_before_tax_is_not_mistaken_for_net_income() -> None:
    """법인세차감전순이익 ends in 순이익; the longer caption must win."""
    assert classify_subtotal("법인세차감전순이익") is SubtotalKind.PROFIT_BEFORE_TAX


# ---------------------------------------------------------------------------
# Locating the statement
# ---------------------------------------------------------------------------


def test_finds_the_statement_past_a_cover_sheet(workbook_path: Path) -> None:
    statement = read_income_statement(workbook_path)

    assert statement.sheet == "손익계산서"
    assert statement.source_file == "fs.xlsx"


def test_detects_the_presentation_scale(workbook_path: Path) -> None:
    """`(단위: 백만원)` means figures are in millions."""
    assert read_income_statement(workbook_path).scale == 6


def test_missing_statement_raises(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "그냥 메모"
    path = tmp_path / "empty.xlsx"
    workbook.save(path)

    with pytest.raises(StatementNotFoundError):
        read_income_statement(path)


def test_unknown_sheet_hint_raises(workbook_path: Path) -> None:
    with pytest.raises(StatementNotFoundError, match="not in the workbook"):
        read_income_statement(workbook_path, options=ReadOptions(sheet="없는시트"))


# ---------------------------------------------------------------------------
# Extraction
# ---------------------------------------------------------------------------


def test_every_row_is_extracted_with_the_right_amount(workbook_path: Path) -> None:
    statement = read_income_statement(workbook_path)
    extracted = {line.raw_label: line.amount for line in statement.lines}

    assert extracted == expected_amounts()


def test_subtotals_are_flagged(workbook_path: Path) -> None:
    statement = read_income_statement(workbook_path)

    flagged = {line.raw_label for line in statement.lines if line.is_subtotal}
    expected = {row.label for row in STATEMENT_ROWS if row.is_subtotal}
    assert flagged == expected


def test_provenance_points_at_the_originating_cell(workbook_path: Path) -> None:
    """Spec §18: the audit trail must reach the exact source cell."""
    statement = read_income_statement(workbook_path)
    revenue = next(line for line in statement.lines if line.raw_label == "매출액")

    assert revenue.locator.sheet == "손익계산서"
    assert revenue.locator.column == "C"
    assert revenue.locator.row == 6
    assert revenue.locator.cell == "C6"
    assert revenue.locator.source_file == "fs.xlsx"


def test_note_references_are_captured(workbook_path: Path) -> None:
    statement = read_income_statement(workbook_path)
    equity_method = next(line for line in statement.lines if line.raw_label == "지분법이익")

    assert equity_method.note_references == ("주석 13",)


def test_indentation_becomes_depth(workbook_path: Path) -> None:
    statement = read_income_statement(workbook_path)
    by_label = {line.raw_label: line.depth for line in statement.lines}

    assert by_label["매출액"] == 0
    assert by_label["기타수익"] == 1


def test_comparative_column_is_read_separately(workbook_path: Path) -> None:
    """Test vector T12: periods must not be mixed."""
    current = read_income_statement(workbook_path)
    prior = read_income_statement(workbook_path, options=ReadOptions(period_index=1))

    current_revenue = next(ln for ln in current.lines if ln.raw_label == "매출액")
    prior_revenue = next(ln for ln in prior.lines if ln.raw_label == "매출액")

    assert current_revenue.amount == Decimal("1000000")
    assert prior_revenue.amount == Decimal("900000")
    assert prior_revenue.locator.column == "D"


def test_requesting_a_missing_period_raises(workbook_path: Path) -> None:
    with pytest.raises(StatementNotFoundError, match="period_index"):
        read_income_statement(workbook_path, options=ReadOptions(period_index=9))


# ---------------------------------------------------------------------------
# Reconciliation against the source's own subtotals (spec §17)
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("style", [SignStyle.PARENTHESES, SignStyle.TRIANGLE])
def test_signed_statements_reconcile_as_read(tmp_path: Path, style: SignStyle) -> None:
    path = build_workbook(tmp_path / f"{style.value}.xlsx", style=style)

    report = reconcile_extraction(read_income_statement(path))

    assert report.passed, [(c.check, str(c.reported), str(c.computed)) for c in report.failures]


def test_parentheses_and_triangle_produce_identical_figures(tmp_path: Path) -> None:
    """The convention is presentation; the extracted numbers must not differ."""
    parens = read_income_statement(build_workbook(tmp_path / "p.xlsx", style=SignStyle.PARENTHESES))
    triangle = read_income_statement(build_workbook(tmp_path / "t.xlsx", style=SignStyle.TRIANGLE))

    assert [ln.amount for ln in parens.lines] == [ln.amount for ln in triangle.lines]


def test_unsigned_statement_fails_as_read_then_reconciles_after_inference(
    tmp_path: Path,
) -> None:
    """The case that matters: figures printed without signs.

    Read literally the statement contradicts its own subtotals. Inference is
    accepted only because reconciliation then passes.
    """
    path = build_workbook(tmp_path / "u.xlsx", style=SignStyle.UNSIGNED)
    statement = read_income_statement(path)

    assert not reconcile_extraction(statement).passed

    adjusted, inferred = infer_signs_from_subtotals(statement)

    assert inferred
    assert reconcile_extraction(adjusted).passed
    assert {ln.raw_label: ln.amount for ln in adjusted.lines} == expected_amounts()


def test_inferred_lines_record_how_their_sign_was_determined(tmp_path: Path) -> None:
    """Spec §18: a derived sign must be traceable, not indistinguishable."""
    path = build_workbook(tmp_path / "u.xlsx", style=SignStyle.UNSIGNED)
    adjusted, _ = infer_signs_from_subtotals(read_income_statement(path))

    cost_of_sales = next(ln for ln in adjusted.lines if ln.raw_label == "매출원가")
    revenue = next(ln for ln in adjusted.lines if ln.raw_label == "매출액")

    assert cost_of_sales.sign_normalization is SignNormalization.INFERRED_FROM_SUBTOTAL
    assert revenue.sign_normalization is SignNormalization.AS_IS


def test_detect_scale_defaults_to_won(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "손익계산서"
    path = tmp_path / "noscale.xlsx"
    workbook.save(path)

    from openpyxl import load_workbook

    loaded = load_workbook(path)
    assert detect_scale(loaded["Sheet"]) == 0
