"""Excel export (spec §34 criteria 11 and 12)."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.adapters.export.xlsx import (
    ExportPackage,
    ExportPrecisionError,
    excel_number,
    export_to_path,
)
from app.adapters.ingest.xlsx import read_income_statement
from app.core.product import DISCLAIMER_EN, DISCLAIMER_KO, UNRECONCILED_WATERMARK
from app.domain.enums import SubtotalKey
from app.domain.impact import analyse
from app.domain.statement import reconstruct
from app.domain.validation import validate
from tests.fixtures.korean_income_statement import SignStyle, build_workbook
from tests.test_reconstruction_end_to_end import (
    NON_FINANCIAL,
    _classify,
    _decompose_aggregates,
)


def _text(sheet: Worksheet, row: int, column: int) -> str:
    value = sheet.cell(row=row, column=column).value
    return "" if value is None else str(value)


def _number(sheet: Worksheet, row: int, column: int) -> float:
    value = sheet.cell(row=row, column=column).value
    assert isinstance(value, int | float), f"expected a number at r{row}c{column}, got {value!r}"
    return float(value)


def _package(tmp_path: Path, *, resolve: bool = True) -> ExportPackage:
    source = read_income_statement(
        build_workbook(tmp_path / "in.xlsx", style=SignStyle.PARENTHESES)
    )
    if resolve:
        source = _decompose_aggregates(source)
    classified = _classify(source, answer_questions=resolve)
    reconstructed = reconstruct(source, classified, facts=NON_FINANCIAL)
    return ExportPackage(
        source=source,
        classified=classified,
        reconstructed=reconstructed,
        validation=validate(source, classified, reconstructed),
        impact=analyse(source, classified, reconstructed),
        project_name="테스트 주식회사",
        rule_set_version="2026.09.1",
        catalog_version="2026.09.1",
    )


@pytest.fixture
def exported(tmp_path: Path) -> Path:
    return export_to_path(_package(tmp_path), tmp_path / "out.xlsx")


# ---------------------------------------------------------------------------
# Precision — the export is the deliverable
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("value", ["1000000", "-70000", "1234.56", "0.1", "95160", "126000.000000"])
def test_realistic_figures_convert_exactly(value: str) -> None:
    assert Decimal(repr(excel_number(Decimal(value)))) == Decimal(value)


@pytest.mark.parametrize(
    "value",
    [
        "12345678901234567890123456789012.123456",  # the full NUMERIC(38,6) range
        "123456789012345678",
    ],
)
def test_a_figure_that_would_lose_a_digit_fails_the_export(value: str) -> None:
    """Shipping a silently-altered number is the one thing not to do.

    openpyxl converts Decimal to a binary float on write, so beyond about
    sixteen significant digits the figure in the file is not the figure on the
    screen. Rather than accept that, the export refuses.
    """
    with pytest.raises(ExportPrecisionError, match="more precision"):
        excel_number(Decimal(value))


def test_exported_figures_equal_the_reconstruction(tmp_path: Path) -> None:
    """The Phase 8 exit criterion: the file says what the screen says."""
    package = _package(tmp_path)
    path = export_to_path(package, tmp_path / "out.xlsx")

    sheet = load_workbook(path)["IFRS 18 Statement"]
    values = {
        _text(sheet, r, 2): sheet.cell(row=r, column=3).value for r in range(1, sheet.max_row + 1)
    }

    assert values["Operating total"] == float(
        package.reconstructed.amount(SubtotalKey.OPERATING_PROFIT)
    )
    assert values["Profit before tax"] == float(
        package.reconstructed.amount(SubtotalKey.PROFIT_BEFORE_TAX)
    )
    assert values["Profit for the period"] == float(
        package.reconstructed.amount(SubtotalKey.PROFIT_FOR_THE_PERIOD)
    )


# ---------------------------------------------------------------------------
# Structure
# ---------------------------------------------------------------------------


def test_every_sheet_is_present(exported: Path) -> None:
    assert load_workbook(exported).sheetnames == [
        "Summary",
        "IFRS 18 Statement",
        "Impact",
        "Accounts",
        "Audit trail",
        "Reconciliation",
    ]


def test_the_disclaimer_is_in_the_file(exported: Path) -> None:
    """Spec §24: the disclaimer travels with the figures, not just the screen."""
    sheet = load_workbook(exported)["Summary"]
    text = " ".join(_text(sheet, r, 1) for r in range(1, sheet.max_row + 1))

    assert DISCLAIMER_EN in text
    assert DISCLAIMER_KO in text
    assert "Does not produce management-defined performance measure" in text


def test_the_bridge_is_exported_and_balances(exported: Path) -> None:
    sheet = load_workbook(exported)["Impact"]
    kinds = {r: _text(sheet, r, 1) for r in range(1, sheet.max_row + 1)}
    start = _number(sheet, next(r for r, k in kinds.items() if k == "START"), 3)
    end = _number(sheet, next(r for r, k in kinds.items() if k == "END"), 3)
    deltas = [_number(sheet, r, 3) for r, k in kinds.items() if k == "DELTA"]

    assert start == 120000
    assert end == 126000
    assert pytest.approx(start + sum(deltas)) == end


def test_every_classification_has_an_audit_record(exported: Path) -> None:
    """Spec §34 criterion 12."""
    sheet = load_workbook(exported)["Audit trail"]
    rows = [[_text(sheet, r, c) for c in range(1, 10)] for r in range(2, sheet.max_row + 1)]

    assert rows
    for account, _category, method, rule, citation, *_rest in rows:
        assert account, "every row names its account"
        assert method, f"{account} has no method"
        if method in {"RULE", "RESIDUAL_DEFAULT"}:
            assert rule, f"{account} was decided by a rule but names none"
            assert citation, f"{account} cites nothing"


def test_accounts_sheet_says_which_rows_sum(exported: Path) -> None:
    """A reader who cannot tell would double count a decomposed caption."""
    sheet = load_workbook(exported)["Accounts"]
    rows = {_text(sheet, r, 1): _text(sheet, r, 3) for r in range(2, sheet.max_row + 1)}

    assert rows["매출액"] == "yes"
    assert rows["금융수익"].startswith("no — decomposed")
    assert rows["이자수익"] == "yes"


def test_included_rows_total_to_profit_for_the_period(exported: Path) -> None:
    sheet = load_workbook(exported)["Accounts"]
    total_row = next(
        r
        for r in range(1, sheet.max_row + 1)
        if _text(sheet, r, 1) == "Total of rows included above"
    )

    assert _number(sheet, total_row, 2) == 95160


def test_the_reconciliation_sheet_lists_every_check(exported: Path) -> None:
    sheet = load_workbook(exported)["Reconciliation"]
    checks = {_text(sheet, r, 1): _text(sheet, r, 2) for r in range(2, sheet.max_row + 1)}

    assert checks["TOTAL_INVARIANCE"] == "PASS"
    assert checks["OPERATING_BRIDGE"] == "PASS"
    assert all(result == "PASS" for result in checks.values())


def test_measures_with_no_before_say_so(exported: Path) -> None:
    """Not a zero, which would claim a change that did not happen."""
    sheet = load_workbook(exported)["Summary"]
    labels = {_text(sheet, r, 1): r for r in range(1, sheet.max_row + 1)}

    assert _text(sheet, labels["PROFIT_BEFORE_FINANCING_AND_INCOME_TAXES"], 2) == "not applicable"
    assert _number(sheet, labels["OPERATING_PROFIT"], 2) == 120000


# ---------------------------------------------------------------------------
# An unreconciled export must never look like a normal one (spec §19)
# ---------------------------------------------------------------------------


def test_an_unreconciled_export_is_watermarked_on_every_sheet(tmp_path: Path) -> None:
    package = _package(tmp_path, resolve=False)
    assert not package.reconciled

    path = export_to_path(package, tmp_path / "blocked.xlsx")

    workbook = load_workbook(path)
    for name in workbook.sheetnames:
        assert workbook[name].cell(row=1, column=1).value == UNRECONCILED_WATERMARK, name


def test_a_reconciled_export_carries_no_watermark(exported: Path) -> None:
    workbook = load_workbook(exported)

    for name in workbook.sheetnames:
        assert workbook[name].cell(row=1, column=1).value != UNRECONCILED_WATERMARK, name


def test_an_unreconciled_export_names_the_failing_check(tmp_path: Path) -> None:
    path = export_to_path(_package(tmp_path, resolve=False), tmp_path / "blocked.xlsx")

    sheet = load_workbook(path)["Reconciliation"]
    failures = {
        _text(sheet, r, 1) for r in range(1, sheet.max_row + 1) if _text(sheet, r, 2) == "FAIL"
    }

    assert failures == {"CATEGORY_COMPLETENESS"}


def test_an_unreconciled_summary_reports_failed(tmp_path: Path) -> None:
    path = export_to_path(_package(tmp_path, resolve=False), tmp_path / "blocked.xlsx")

    sheet = load_workbook(path)["Summary"]
    rows = {_text(sheet, r, 1): _text(sheet, r, 2) for r in range(1, sheet.max_row + 1)}

    assert rows["Reconciliation"] == "FAILED"
